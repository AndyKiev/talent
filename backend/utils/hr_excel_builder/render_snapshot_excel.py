"""
render_snapshot_excel.py
────────────────────────
Renders a staffing-snapshot Excel report from a v2 snapshot JSON file.

Usage:
    python render_snapshot_excel.py snapshot_v2.json output.xlsx [ukr|eng]

Header labels are localized: pass a ``labels`` dict to ``render_excel`` (see
``snapshot_labels`` — resolves the hrSnapshot* message keys from the DB for the
target language), or give the CLI a language code (default ukr). The in-code
``DEFAULT_LABELS`` fallback is English, per the project translation rule.

The JSON schema is config-driven: each job_group carries a `config` block
that controls both data shape and Excel column layout:

    target_mode : "total"            → single target column  (no Pa/Po split)
                  "by_status"        → two target columns     (Па | По)

    fact_mode   : "by_status"        → two fact columns       (Па | По)
                  "by_job_and_status" → 2 × len(jobs) columns (one pair per job)

    jobs        : []                 → no sub-job row-2 header
                  [{id, key, name}]  → adds row-2 sub-job breakdown header

Column layout is computed dynamically from the config, so adding/removing a
job_group or changing its modes requires no changes to this renderer.

Dependencies:
    pip install openpyxl
"""

import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────────────────────
# Header labels (localized; English = in-code fallback)
# ─────────────────────────────────────────────────────────────────────────────
DEFAULT_LABELS = {
    "object": "Object",
    "total": "Total",
    "target": "target",
    "fact": "actual",
    "pct": "% done",
    "base_target": "base target",
    "object_target": "object target",
    "status_pa": "Pas",
    "status_po": "Cur",
}

# slot -> msg_keys.name (resolved from the DB by snapshot_labels)
LABEL_KEYS = {
    "object": "hrSnapshotObject",
    "total": "hrSnapshotTotal",
    "target": "hrSnapshotTarget",
    "fact": "hrSnapshotFact",
    "pct": "hrSnapshotPctDone",
    "base_target": "hrSnapshotBaseTarget",
    "object_target": "hrSnapshotObjectTarget",
    "status_pa": "hrSnapshotStatusPa",
    "status_po": "hrSnapshotStatusPo",
}


async def snapshot_labels(session, lang_id: int) -> dict:
    """Resolve the header labels from the msg_keys/msgs tables for one language.

    Missing keys keep their English DEFAULT_LABELS value. Imports are lazy so
    the renderer stays importable/runnable without the backend DB stack.
    """
    from backend.api_v1.msg_pg.msg_translate import translate_keys

    key_to_slot = {key: slot for slot, key in LABEL_KEYS.items()}
    keys = {LABEL_KEYS[slot]: fallback for slot, fallback in DEFAULT_LABELS.items()}
    resolved = await translate_keys(session, keys, lang_id)
    return {key_to_slot[key]: value for key, value in resolved.items()}


# ─────────────────────────────────────────────────────────────────────────────
# Colour palette
# ─────────────────────────────────────────────────────────────────────────────
C_H1 = "1F3864"  # dark navy   – job-group row 1 header
C_H2 = "2E75B6"  # mid blue    – sub-job row 2 header
C_H3 = "BDD7EE"  # light blue  – target/fact row 3 header
C_H4 = "DEEAF1"  # very light  – Па/По row 4 header
C_TOT = "FFF2CC"  # yellow      – grand-totals row
C_PCT = "E2EFDA"  # green tint  – % columns
C_ODD = "FFFFFF"  # white       – odd data rows
C_EVEN = "EBF3FB"  # pale blue   – even data rows


# ─────────────────────────────────────────────────────────────────────────────
# Style helpers
# ─────────────────────────────────────────────────────────────────────────────
def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def _thin_border() -> Border:
    s = Side(style="thin", color="BFBFBF")
    return Border(left=s, right=s, top=s, bottom=s)


def _hdr_font(color: str = "FFFFFF", bold: bool = True, size: int = 9) -> Font:
    return Font(bold=bold, color=color, size=size, name="Arial")


def _data_font(bold: bool = False, size: int = 9) -> Font:
    return Font(bold=bold, size=size, name="Arial")


ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _set_cell(
    ws,
    row: int,
    col: int,
    value,
    font=None,
    fill_color: str = None,
    align=None,
    num_fmt: str = None,
):
    cell = ws.cell(row, col, value)
    if font:
        cell.font = font
    if fill_color:
        cell.fill = _fill(fill_color)
    if align:
        cell.alignment = align
    if num_fmt:
        cell.number_format = num_fmt
    cell.border = _thin_border()
    return cell


def _merge(
    ws, r1: int, c1: int, r2: int, c2: int, value, font=None, fill_color: str = None
):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    cell = ws.cell(r1, c1, value)
    if font:
        cell.font = font
    if fill_color:
        cell.fill = _fill(fill_color)
    cell.alignment = ALIGN_CENTER
    cell.border = _thin_border()


# ─────────────────────────────────────────────────────────────────────────────
# Column-plan builder
# ─────────────────────────────────────────────────────────────────────────────
def _build_column_plan(job_groups: list) -> tuple:
    """
    Walk job_groups and assign concrete column indices based on each group's
    config.  Returns (jg_plans, summary_cols, max_col).

    jg_plans  – list of dicts, one per job_group, with keys:
        jg            – the original job_group dict
        col_start     – first column index of this group
        col_end       – last column index of this group
        target_cols   – list of 1 or 2 column indices for target
        fact_cols     – list of 2 column indices (by_status only)
        job_fact_cols – {job_key: [pa_col, po_col]}  (by_job_and_status only)
        pct_col       – column index for the % column

    summary_cols – dict with keys: base_target, object_target, fact, pct
    max_col      – total number of columns
    """
    COL_ORG = 1
    COL_DEPT = 2
    cur = 3

    jg_plans = []
    for jg in job_groups:
        cfg = jg["config"]
        plan = {"jg": jg, "col_start": cur}

        # Target columns
        if cfg["target_mode"] == "total":
            plan["target_cols"] = [cur]
            cur += 1
        else:
            plan["target_cols"] = [cur, cur + 1]
            cur += 2

        # Fact columns
        if cfg["fact_mode"] == "by_status":
            plan["fact_cols"] = [cur, cur + 1]
            cur += 2
            plan["job_fact_cols"] = {}
        else:  # by_job_and_status
            plan["fact_cols"] = []
            plan["job_fact_cols"] = {}
            for job in cfg["jobs"]:
                plan["job_fact_cols"][job["key"]] = [cur, cur + 1]
                cur += 2

        plan["pct_col"] = cur
        cur += 1
        plan["col_end"] = cur - 1
        jg_plans.append(plan)

    summary_cols = {
        "base_target": cur,
        "object_target": cur + 1,
        "fact": cur + 2,
        "pct": cur + 3,
    }
    max_col = cur + 3
    return jg_plans, summary_cols, max_col


# ─────────────────────────────────────────────────────────────────────────────
# Header writers  (rows 1 – 4)
# ─────────────────────────────────────────────────────────────────────────────
def _write_headers(ws, jg_plans: list, summary_cols: dict, labels: dict):
    sc = summary_cols
    COL_ORG, COL_DEPT = 1, 2

    # ── Row 1 & 2: job-group titles ───────────────────────────────────────────
    # The object column always spans rows 1–2
    _merge(
        ws, 1, COL_ORG, 2, COL_DEPT, labels["object"], font=_hdr_font(), fill_color=C_H1
    )

    for p in jg_plans:
        has_sub_jobs = bool(p["jg"]["config"]["jobs"])
        row2_end = 1 if has_sub_jobs else 2  # span both rows when no sub-job row
        _merge(
            ws,
            1,
            p["col_start"],
            row2_end,
            p["col_end"],
            p["jg"]["key"],
            font=_hdr_font(),
            fill_color=C_H1,
        )

    # The totals block always spans rows 1–2
    _merge(
        ws,
        1,
        sc["base_target"],
        2,
        sc["pct"],
        labels["total"],
        font=_hdr_font(),
        fill_color=C_H1,
    )

    # ── Row 2: sub-job labels (only for groups that have jobs) ────────────────
    for p in jg_plans:
        cfg = p["jg"]["config"]
        if not cfg["jobs"]:
            continue
        # Target label in row 2
        if cfg["target_mode"] == "by_status":
            _merge(
                ws,
                2,
                p["target_cols"][0],
                2,
                p["target_cols"][-1],
                labels["target"],
                font=_hdr_font(),
                fill_color=C_H2,
            )
        else:
            _set_cell(
                ws,
                2,
                p["target_cols"][0],
                labels["target"],
                font=_hdr_font(),
                fill_color=C_H2,
                align=ALIGN_CENTER,
            )
        # Per-job fact labels
        for job in cfg["jobs"]:
            jc = p["job_fact_cols"][job["key"]]
            _merge(
                ws, 2, jc[0], 2, jc[1], job["key"], font=_hdr_font(), fill_color=C_H2
            )
        # % label
        _set_cell(
            ws,
            2,
            p["pct_col"],
            "%",
            font=_hdr_font(),
            fill_color=C_H2,
            align=ALIGN_CENTER,
        )

    # ── Row 3 (A3:B4 label block) + target/fact/% row labels ─────────────────
    _merge(ws, 3, COL_ORG, 4, COL_DEPT, "", font=_hdr_font("1F3864"), fill_color=C_H3)

    for p in jg_plans:
        cfg = p["jg"]["config"]

        # Target label
        if cfg["target_mode"] == "total":
            _set_cell(
                ws,
                3,
                p["target_cols"][0],
                labels["target"],
                font=_hdr_font("1F3864", bold=False),
                fill_color=C_H3,
                align=ALIGN_CENTER,
            )
        else:
            _merge(
                ws,
                3,
                p["target_cols"][0],
                3,
                p["target_cols"][1],
                labels["target"],
                font=_hdr_font("1F3864", bold=False),
                fill_color=C_H3,
            )

        # Fact label spanning all fact columns
        if cfg["fact_mode"] == "by_status":
            _merge(
                ws,
                3,
                p["fact_cols"][0],
                3,
                p["fact_cols"][1],
                labels["fact"],
                font=_hdr_font("1F3864", bold=False),
                fill_color=C_H3,
            )
        else:
            all_fact = [c for jc in p["job_fact_cols"].values() for c in jc]
            _merge(
                ws,
                3,
                min(all_fact),
                3,
                max(all_fact),
                labels["fact"],
                font=_hdr_font("1F3864", bold=False),
                fill_color=C_H3,
            )

        # % done
        _set_cell(
            ws,
            3,
            p["pct_col"],
            labels["pct"],
            font=_hdr_font("1F3864", bold=False),
            fill_color=C_PCT,
            align=ALIGN_CENTER,
        )

    # Summary row-3 labels
    for col, label in [
        (sc["base_target"], labels["base_target"]),
        (sc["object_target"], labels["object_target"]),
        (sc["fact"], labels["fact"]),
        (sc["pct"], labels["pct"]),
    ]:
        fc = C_PCT if col == sc["pct"] else C_H3
        _set_cell(
            ws,
            3,
            col,
            label,
            font=_hdr_font("1F3864", bold=False),
            fill_color=fc,
            align=ALIGN_CENTER,
        )

    # ── Row 4: Па / По labels ─────────────────────────────────────────────────
    for p in jg_plans:
        cfg = p["jg"]["config"]

        if cfg["target_mode"] == "by_status":
            _set_cell(
                ws,
                4,
                p["target_cols"][0],
                labels["status_pa"],
                font=_hdr_font("1F3864", bold=False, size=8),
                fill_color=C_H4,
                align=ALIGN_CENTER,
            )
            _set_cell(
                ws,
                4,
                p["target_cols"][1],
                labels["status_po"],
                font=_hdr_font("1F3864", bold=False, size=8),
                fill_color=C_H4,
                align=ALIGN_CENTER,
            )
        else:
            _set_cell(
                ws,
                4,
                p["target_cols"][0],
                "",
                font=_hdr_font("1F3864", bold=False, size=8),
                fill_color=C_H4,
                align=ALIGN_CENTER,
            )

        if cfg["fact_mode"] == "by_status":
            _set_cell(
                ws,
                4,
                p["fact_cols"][0],
                labels["status_pa"],
                font=_hdr_font("1F3864", bold=False, size=8),
                fill_color=C_H4,
                align=ALIGN_CENTER,
            )
            _set_cell(
                ws,
                4,
                p["fact_cols"][1],
                labels["status_po"],
                font=_hdr_font("1F3864", bold=False, size=8),
                fill_color=C_H4,
                align=ALIGN_CENTER,
            )
        else:
            for jc in p["job_fact_cols"].values():
                _set_cell(
                    ws,
                    4,
                    jc[0],
                    labels["status_pa"],
                    font=_hdr_font("1F3864", bold=False, size=8),
                    fill_color=C_H4,
                    align=ALIGN_CENTER,
                )
                _set_cell(
                    ws,
                    4,
                    jc[1],
                    labels["status_po"],
                    font=_hdr_font("1F3864", bold=False, size=8),
                    fill_color=C_H4,
                    align=ALIGN_CENTER,
                )

        _set_cell(
            ws,
            4,
            p["pct_col"],
            "",
            font=_hdr_font("1F3864", bold=False),
            fill_color=C_PCT,
            align=ALIGN_CENTER,
        )

    for col in [sc["base_target"], sc["object_target"], sc["fact"]]:
        _set_cell(
            ws,
            4,
            col,
            "",
            font=_hdr_font("1F3864"),
            fill_color=C_H4,
            align=ALIGN_CENTER,
        )
    _set_cell(
        ws,
        4,
        sc["pct"],
        "",
        font=_hdr_font("1F3864"),
        fill_color=C_PCT,
        align=ALIGN_CENTER,
    )


# ─────────────────────────────────────────────────────────────────────────────
# Data-row writer  (grand-totals row + each store row)
# ─────────────────────────────────────────────────────────────────────────────
def _write_data_row(
    ws,
    excel_row: int,
    row_data: dict,
    jg_plans: list,
    summary_cols: dict,
    labels: dict,
    is_total: bool = False,
    row_index: int = 0,
):
    """
    Write one data row.  `row_index` (0-based) controls alternating row colour.
    """
    COL_ORG, COL_DEPT = 1, 2
    sc = summary_cols

    def _bg(col: int) -> str:
        if is_total:
            return C_TOT
        return C_ODD if row_index % 2 == 0 else C_EVEN

    def _put(col: int, value, is_pct: bool = False, bold: bool = False):
        fmt = "0%" if is_pct else "0"
        fc = C_PCT if is_pct else _bg(col)
        val = value if value is not None else 0
        font = _hdr_font("7F6000", bold=bold) if is_total else _data_font(bold)
        _set_cell(
            ws,
            excel_row,
            col,
            val,
            font=font,
            fill_color=fc,
            align=ALIGN_CENTER,
            num_fmt=fmt,
        )

    # Org-unit / department columns
    if is_total:
        _merge(
            ws,
            excel_row,
            COL_ORG,
            excel_row,
            COL_DEPT,
            row_data.get("label") or labels["total"],
            font=_hdr_font("7F6000", bold=True),
            fill_color=C_TOT,
        )
    else:
        _set_cell(
            ws,
            excel_row,
            COL_ORG,
            row_data["org_unit_key"],
            font=_data_font(bold=True),
            fill_color=_bg(COL_ORG),
            align=ALIGN_LEFT,
        )
        _set_cell(
            ws,
            excel_row,
            COL_DEPT,
            row_data["department_key"],
            font=_data_font(),
            fill_color=_bg(COL_DEPT),
            align=ALIGN_CENTER,
        )

    # Job-group columns
    for p in jg_plans:
        jg_id = p["jg"]["id"]
        cfg = p["jg"]["config"]
        jg_data = row_data.get("job_groups", {}).get(jg_id, {})

        # Target
        if cfg["target_mode"] == "total":
            _put(p["target_cols"][0], jg_data.get("target"))
        else:
            t = jg_data.get("target") or {}
            _put(p["target_cols"][0], t.get("pa"))
            _put(p["target_cols"][1], t.get("po"))

        # Fact
        if cfg["fact_mode"] == "by_status":
            f = jg_data.get("fact") or {}
            _put(p["fact_cols"][0], f.get("pa"))
            _put(p["fact_cols"][1], f.get("po"))
        else:
            for job in cfg["jobs"]:
                jf = (jg_data.get("fact") or {}).get(job["key"]) or {}
                jc = p["job_fact_cols"][job["key"]]
                _put(jc[0], jf.get("pa"))
                _put(jc[1], jf.get("po"))

        # %
        _put(p["pct_col"], jg_data.get("pct"), is_pct=True)

    # Summary columns
    s = row_data.get("summary") or {}
    _put(sc["base_target"], s.get("base_target"))
    _put(sc["object_target"], s.get("object_target"))
    _put(sc["fact"], s.get("fact"))
    _put(sc["pct"], s.get("pct"), is_pct=True)


# ─────────────────────────────────────────────────────────────────────────────
# Column widths & row heights
# ─────────────────────────────────────────────────────────────────────────────
def _apply_dimensions(
    ws, jg_plans: list, summary_cols: dict, max_col: int, data_row_count: int
):
    ws.column_dimensions[get_column_letter(1)].width = 11  # org_unit
    ws.column_dimensions[get_column_letter(2)].width = 9  # department

    pct_cols = {p["pct_col"] for p in jg_plans} | {summary_cols["pct"]}
    for col in range(3, max_col + 1):
        ws.column_dimensions[get_column_letter(col)].width = (
            7.0 if col in pct_cols else 5.5
        )

    for row in range(1, 5):
        ws.row_dimensions[row].height = 30
    ws.row_dimensions[5].height = 18
    for row in range(6, 6 + data_row_count):
        ws.row_dimensions[row].height = 15


# ─────────────────────────────────────────────────────────────────────────────
# Public entry point
# ─────────────────────────────────────────────────────────────────────────────
def render_excel(snapshot: dict, out_path: str, labels: dict = None) -> None:
    """
    Render a v2 staffing snapshot dict to an Excel file at `out_path`.

    Parameters
    ----------
    snapshot  : dict loaded from snapshot_v2.json
    out_path  : destination .xlsx path
    labels    : header labels for the target language (see ``snapshot_labels``);
                missing slots fall back to the English DEFAULT_LABELS
    """
    labels = {**DEFAULT_LABELS, **(labels or {})}
    job_groups = snapshot["essences"]["job_groups"]
    jg_plans, summary_cols, max_col = _build_column_plan(job_groups)

    wb = Workbook()
    ws = wb.active
    ws.title = "Snapshot"

    _write_headers(ws, jg_plans, summary_cols, labels)

    # Grand-totals row → always row 5
    _write_data_row(
        ws, 5, snapshot["grand_totals"], jg_plans, summary_cols, labels, is_total=True
    )

    # Data rows → rows 6 +
    for i, row in enumerate(snapshot["data"]):
        _write_data_row(ws, 6 + i, row, jg_plans, summary_cols, labels, row_index=i)

    _apply_dimensions(
        ws, jg_plans, summary_cols, max_col, data_row_count=len(snapshot["data"])
    )

    # Freeze header rows and first two identifier columns
    ws.freeze_panes = f"{get_column_letter(3)}6"

    wb.save(out_path)
    print(f"Saved: {out_path}")


# ─────────────────────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────────────────────
def _cli_fetch_labels(lang_code: str) -> dict:
    """Resolve labels from the DB for the CLI (ukr/eng); English fallback when
    the backend stack or the DB is unavailable."""
    import asyncio
    from pathlib import Path

    repo_root = Path(__file__).resolve().parents[3]
    if str(repo_root) not in sys.path:
        sys.path.insert(0, str(repo_root))
    try:
        from backend.api_v1.msg_pg.msg_translate import LANG_ID_ENG, LANG_ID_UKR
        from backend.database.db_helper import db_helper

        lang_id = LANG_ID_UKR if lang_code == "ukr" else LANG_ID_ENG

        async def _fetch() -> dict:
            async with db_helper.session_factory() as session:
                return await snapshot_labels(session, lang_id)

        return asyncio.run(_fetch())
    except Exception as exc:
        print(f"[WARN] Could not load '{lang_code}' labels from the DB ({exc}); "
              "using English fallbacks.")
        return dict(DEFAULT_LABELS)


if __name__ == "__main__":
    if len(sys.argv) not in (3, 4):
        print(
            "Usage: python render_snapshot_excel.py <snapshot.json> <output.xlsx> [ukr|eng]"
        )
        sys.exit(1)

    json_path, xlsx_path = sys.argv[1], sys.argv[2]
    lang_code = sys.argv[3] if len(sys.argv) == 4 else "ukr"
    if lang_code not in ("ukr", "eng"):
        print(f"Unknown language '{lang_code}' — use ukr or eng.")
        sys.exit(1)

    with open(json_path, encoding="utf-8") as f:
        data = json.load(f)

    render_excel(data, xlsx_path, labels=_cli_fetch_labels(lang_code))
