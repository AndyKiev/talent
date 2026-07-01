from typing import List, Optional

from sqlalchemy.exc import IntegrityError
from sqlalchemy.ext.asyncio import AsyncSession

from backend.api_v1.base.base_service import BaseService
from backend.api_v1.base.mutation_response import MutationResponse
from backend.api_v1.job.job_repository import JobRepository
from backend.api_v1.job.job_schema import Job as JobSchema, JobCreate, JobUpdate
from backend.api_v1.employee.employee_schema import EmployeeSchema as UserSchema
from backend.api_v1.job.job_errors import (
    JobNotFound,
    JobNameTaken,
    JobDeleteError,
    JobNotFoundByName,
)
from backend.api_v1.job.job_success import (
    JobDeleteSuccess,
    JobCreateSuccess,
    JobUpdateSuccess,
)
from backend.api_v1.employee.employee_repository import EmployeeRepository
from backend.api_v1.employee.employee_service import EmployeeService, SyncJobResult
from backend.api_v1.base.errors import DomainError

import io
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from fastapi import UploadFile
from backend.api_v1.job.job_schema import JobBulkRow, JobBulkUploadResult
from backend.api_v1.job.job_errors import (
    JobBulkUploadNothingToInsert,
    JobBulkUploadInvalidFile,
)
from backend.api_v1.job.job_success import JobBulkUploadSuccess


class JobService(BaseService):
    def __init__(
        self,
        repository: JobRepository,
        user: Optional[UserSchema] = None,
        session: Optional[AsyncSession] = None,
    ):
        super().__init__(repository, user=user, session=session)

    # ------------------------------------------------------------------
    # Read
    # ------------------------------------------------------------------

    async def get_by_id(self, id: int) -> JobSchema:
        result = await self.repository.get_by_id(id)
        if not result:
            exc = JobNotFound(id)
            raise await self._resolve_domain_error(exc)
        return result

    async def get_jobs(
        self,
        name: Optional[str] = None,
        sort: Optional[str] = None,
    ) -> List[JobSchema]:
        if name:
            job = await self.get_by_name(name, not_found_exc=JobNotFoundByName)
            return [JobSchema.model_validate(job)]
        jobs = await self.repository.get_all_with_dept_type_links()
        return [JobSchema.model_validate(j) for j in jobs]

    # ------------------------------------------------------------------
    # Write
    # ------------------------------------------------------------------

    async def create_job(self, job_in: JobCreate) -> MutationResponse[JobSchema]:
        await self.exists_by_name(job_in.name, already_exists_exc=JobNameTaken)
        try:
            job = await self.create(job_in)
            # Optional-essence-property: auto-link the default category when the
            # app-level flag is on. Reload so the response reflects the link.
            if await self._apply_category_on_create_enabled():
                await self._link_default_category(job.id)
                job = await self.repository.get_by_id(job.id)
            schema = JobSchema.model_validate(job)
            detail = await self._resolve_domain_success(JobCreateSuccess(schema.name))
            return MutationResponse(detail=detail, data=schema)
        except IntegrityError:
            raise await self._resolve_domain_error(JobNameTaken(job_in.name))

    # ------------------------------------------------------------------
    # Optional category-on-create (see /optional-essence-property)
    # ------------------------------------------------------------------

    async def _apply_category_on_create_enabled(self) -> bool:
        from backend.api_v1.app_setting.app_setting_service import get_bool_setting

        return await get_bool_setting(
            self.session, "job_apply_category_on_create", default=True
        )

    async def _default_category_id(self) -> Optional[int]:
        """Resolve the default ('manager') category by KEY — survives reseed/id reorder."""
        from sqlalchemy import select
        from backend.api_v1.job_category.job_category_model import JobCategory

        return await self.session.scalar(
            select(JobCategory.id).where(JobCategory.key == "manager")
        )

    async def _link_default_category(self, job_id: int) -> None:
        from sqlalchemy import select
        from backend.api_v1.job_job_category_link.job_job_category_link_model import (
            JobJobCategoryLink,
        )

        category_id = await self._default_category_id()
        if category_id is None:
            return
        existing = await self.session.scalar(
            select(JobJobCategoryLink.id).where(JobJobCategoryLink.job_id == job_id)
        )
        if existing is not None:
            return
        self.session.add(
            JobJobCategoryLink(job_id=job_id, job_category_id=category_id)
        )
        await self.session.commit()

    async def update_job(
        self, job_id: int, job_update: JobUpdate
    ) -> MutationResponse[JobSchema]:
        if job_update.name:
            await self.exists_by_name(job_update.name, already_exists_exc=JobNameTaken)
        try:
            orm_job = await self.get_by_id(job_id)
            updated = await self.update(orm_job, job_update, partial=True)
            schema = JobSchema.model_validate(updated)
            detail = await self._resolve_domain_success(JobUpdateSuccess(schema.name))
            return MutationResponse(detail=detail, data=schema)
        except IntegrityError:
            raise await self._resolve_domain_error(JobNameTaken(job_update.name))

    async def delete_job(self, job_id: int) -> None:
        job = await self.get_by_id(job_id)
        await self.delete_by_id(
            job_id,
            name=job.name,
            delete_error_exc=JobDeleteError,
            delete_success_exc=JobDeleteSuccess,
        )

    # ------------------------------------------------------------------
    # Group management
    # ------------------------------------------------------------------

    async def add_to_group(self, job_id: int, user_group_id: int) -> JobSchema:
        try:
            job = await self.repository.add_to_group(job_id, user_group_id)
            return JobSchema.model_validate(job)
        except DomainError as exc:
            raise await self._resolve_domain_error(exc)

    async def remove_from_group(self, job_id: int, user_group_id: int) -> JobSchema:
        try:
            job = await self.repository.remove_from_group(job_id, user_group_id)
            return JobSchema.model_validate(job)
        except DomainError as exc:
            raise await self._resolve_domain_error(exc)

    async def set_groups(self, job_id: int, user_group_ids: List[int]) -> JobSchema:
        try:
            job = await self.repository.set_groups(job_id, user_group_ids)
            return JobSchema.model_validate(job)
        except DomainError as exc:
            raise await self._resolve_domain_error(exc)

    # ------------------------------------------------------------------
    # Job-group sync
    # ------------------------------------------------------------------

    async def sync_job_users_groups(self, job_id: int) -> "SyncJobResult":
        user_service = EmployeeService(
            repository=EmployeeRepository(session=self.session),
            user=self.user,
            session=self.session,
        )
        return await user_service.sync_job_users_groups(job_id)

    async def bulk_upload_jobs(self, file: "UploadFile") -> "JobBulkUploadResult":
        """
        Parse an Excel file with columns [name, description], filter out rows
        whose name or description already exist in the DB, insert the rest.

        Skipping rules
        --------------
        A row is skipped if **either** condition holds:
          - A job with the same ``name`` (case-insensitive strip) already exists.
          - A job with the same ``description`` (case-insensitive strip) already exists
            AND that description is non-empty.

        Returns
        -------
        JobBulkUploadResult with inserted jobs + skip details.
        Raises JobBulkUploadInvalidFile (→ 422) when the file cannot be parsed.
        Raises JobBulkUploadNothingToInsert (→ 409) when every row is a duplicate.
        """
        # ── 1. Parse the Excel file ───────────────────────────────────────────
        raw = await file.read()
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        except Exception:
            raise await self._resolve_domain_error(
                JobBulkUploadInvalidFile("Cannot open file as an Excel workbook.")
            )

        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)

        # Expect first row to be a header containing "name" and "description"
        try:
            header = [
                str(c).strip().lower() if c is not None else "" for c in next(rows_iter)
            ]
        except StopIteration:
            raise await self._resolve_domain_error(
                JobBulkUploadInvalidFile("File is empty.")
            )

        if "name" not in header:
            raise await self._resolve_domain_error(
                JobBulkUploadInvalidFile(
                    f"Missing required column 'name'. Found columns: {header}"
                )
            )

        name_idx = header.index("name")
        desc_idx = header.index("description") if "description" in header else None

        # Collect valid rows from the file
        file_rows: list[JobBulkRow] = []
        for raw_row in rows_iter:
            name_val = raw_row[name_idx] if name_idx < len(raw_row) else None
            if name_val is None:
                continue
            name_str = str(name_val).strip()
            if not name_str:
                continue
            desc_str: str | None = None
            if desc_idx is not None:
                d = raw_row[desc_idx] if desc_idx < len(raw_row) else None
                desc_str = str(d).strip() if d is not None else None

            file_rows.append(JobBulkRow(name=name_str, description=desc_str or None))

        if not file_rows:
            raise await self._resolve_domain_error(
                JobBulkUploadInvalidFile("No data rows found after the header.")
            )

        # ── 2. Fetch existing jobs from the DB ────────────────────────────────
        existing_jobs = await self.get_all()  # returns list of ORM Job objects
        existing_names: set[str] = {j.name.strip().lower() for j in existing_jobs}
        existing_descs: set[str] = {
            j.description.strip().lower()
            for j in existing_jobs
            if j.description and j.description.strip()
        }

        # ── 3. Partition: to-insert vs. skipped ──────────────────────────────
        to_insert: list[JobBulkRow] = []
        skipped_by_name: list[str] = []
        skipped_by_desc: list[str] = []

        for row in file_rows:
            if row.name.lower() in existing_names:
                skipped_by_name.append(row.name)
                continue
            if row.description and row.description.strip().lower() in existing_descs:
                skipped_by_desc.append(row.name)
                continue
            to_insert.append(row)

        if not to_insert:
            raise await self._resolve_domain_error(
                JobBulkUploadNothingToInsert(skipped=len(file_rows))
            )

        # ── 4. Insert the new jobs ────────────────────────────────────────────
        # Optional-essence-property: bulk-created jobs follow the same gate.
        # Read the flag + default category ONCE before the loop.
        from backend.api_v1.job_job_category_link.job_job_category_link_model import (
            JobJobCategoryLink,
        )

        default_category_id: int | None = None
        if await self._apply_category_on_create_enabled():
            default_category_id = await self._default_category_id()

        inserted_schemas: list[JobSchema] = []
        for row in to_insert:
            job_create = JobCreate(
                name=row.name,
                description=row.description,
                is_active=True,
            )
            orm_job = await self.create(job_create)
            if default_category_id is not None:
                self.session.add(
                    JobJobCategoryLink(
                        job_id=orm_job.id, job_category_id=default_category_id
                    )
                )
                await self.session.commit()
            inserted_schemas.append(JobSchema.model_validate(orm_job))

        # ── 5. Build response ─────────────────────────────────────────────────
        skipped_total = len(skipped_by_name) + len(skipped_by_desc)
        detail = await self._resolve_domain_success(
            JobBulkUploadSuccess(
                inserted=len(inserted_schemas),
                skipped=skipped_total,
            )
        )

        return JobBulkUploadResult(
            detail=detail,
            inserted=inserted_schemas,
            skipped_names=skipped_by_name,
            skipped_descriptions=skipped_by_desc,
            inserted_count=len(inserted_schemas),
            skipped_count=skipped_total,
        )
